from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

navy = RGBColor(18, 52, 86)
blue = RGBColor(45, 115, 188)
light = RGBColor(239, 247, 255)
text = RGBColor(30, 30, 30)
muted = RGBColor(86, 96, 106)
white = RGBColor(255, 255, 255)


def add_paragraphs(tf, lines, font_size=16, color=text, bold=False, space_after=0):
    tf.clear()
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.space_after = Pt(space_after)
        p.level = 0


def add_title(slide, text_value):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12.0), Inches(0.6))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text_value
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = navy
    return box


# Slide 1 - Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = light

bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.7))
bar.fill.solid(); bar.fill.fore_color.rgb = navy
bar.line.fill.background()

header = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(0.8))
header_tf = header.text_frame
p = header_tf.paragraphs[0]
p.text = 'Supply Chain Risk Insight: From Business Problem to Solution'
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = navy

sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(12.0), Inches(0.5))
sub_tf = sub.text_frame
p = sub_tf.paragraphs[0]
p.text = 'FMN Supply Chain AI Solution Walkthrough'
p.font.size = Pt(18)
p.font.color.rgb = muted

# Problem box
problem_box = slide.shapes.add_shape(1, Inches(0.8), Inches(2.5), Inches(4.0), Inches(2.9))
problem_box.fill.solid(); problem_box.fill.fore_color.rgb = white
problem_box.line.color.rgb = blue
problem_box.line.width = Pt(2)
problem_tf = problem_box.text_frame
problem_tf.word_wrap = True
problem_tf.margin_left = Inches(0.2)
problem_tf.margin_right = Inches(0.2)
add_paragraphs(
    problem_tf,
    [
        'Business Problem',
        '• Stockouts and excess inventory were reducing service and increasing cost',
        '• Demand variability made forecasting inconsistent across SKUs',
        '• Teams needed prioritization based on real business risk',
    ],
    font_size=15,
    color=text,
    bold=False,
    space_after=8,
)
problem_tf.paragraphs[0].font.bold = True
problem_tf.paragraphs[0].font.size = Pt(20)
problem_tf.paragraphs[0].font.color.rgb = navy

# Solution box
solution_box = slide.shapes.add_shape(1, Inches(5.2), Inches(2.5), Inches(4.0), Inches(2.9))
solution_box.fill.solid(); solution_box.fill.fore_color.rgb = white
solution_box.line.color.rgb = blue
solution_box.line.width = Pt(2)
solution_tf = solution_box.text_frame
solution_tf.word_wrap = True
solution_tf.margin_left = Inches(0.2)
solution_tf.margin_right = Inches(0.2)
add_paragraphs(
    solution_tf,
    [
        'Solution',
        '• Forecasted demand using historical sales and operational signals',
        '• Flagged SKU risk using cover ratio, lead time, and trend indicators',
        '• Delivered explainable insights and Q&A for decision support',
    ],
    font_size=15,
    color=text,
    bold=False,
    space_after=8,
)
solution_tf.paragraphs[0].font.bold = True
solution_tf.paragraphs[0].font.size = Pt(20)
solution_tf.paragraphs[0].font.color.rgb = navy

# Business value box
value_box = slide.shapes.add_shape(1, Inches(9.6), Inches(2.5), Inches(3.0), Inches(2.9))
value_box.fill.solid(); value_box.fill.fore_color.rgb = navy
value_box.line.color.rgb = navy
value_tf = value_box.text_frame
value_tf.word_wrap = True
value_tf.margin_left = Inches(0.2)
value_tf.margin_right = Inches(0.2)
add_paragraphs(
    value_tf,
    ['Business Value', 'Reduced risk, better planning, faster inventory decisions'],
    font_size=14,
    color=white,
    bold=False,
    space_after=8,
)
value_tf.paragraphs[0].font.bold = True
value_tf.paragraphs[0].font.size = Pt(17)

footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.8), Inches(12.0), Inches(0.4))
ft = footer.text_frame
p = ft.paragraphs[0]
p.text = 'AI-powered supply chain risk management for SKU prioritization and replenishment visibility'
p.font.size = Pt(11)
p.font.color.rgb = muted

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = white
add_title(slide, '1. Business Problem')

problem = slide.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(11.5), Inches(4.8))
ptf = problem.text_frame
ptf.word_wrap = True
items = [
    '• Inventory teams were making decisions without a clear risk signal for every SKU.',
    '• This created both stockouts and excess inventory, hurting service levels and cash flow.',
    '• The challenge was to move from reactive replenishment to proactive, data-driven stock risk management.',
]
add_paragraphs(ptf, items, font_size=22, color=text, space_after=14)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = light
add_title(slide, '2. Solution Overview')

box1 = slide.shapes.add_shape(1, Inches(0.8), Inches(1.5), Inches(3.5), Inches(4.5))
box1.fill.solid(); box1.fill.fore_color.rgb = white
box1.line.color.rgb = blue
box1.line.width = Pt(2)
box1_tf = box1.text_frame
box1_tf.word_wrap = True
add_paragraphs(
    box1_tf,
    [
        'Data + Features',
        'Historic sales, lead time, stock levels, and trend indicators were used to build demand features.',
        'This captured both volume patterns and inventory risk drivers.',
    ],
    font_size=17,
    color=text,
    space_after=8,
)
box1_tf.paragraphs[0].font.bold = True
box1_tf.paragraphs[0].font.size = Pt(22)
box1_tf.paragraphs[0].font.color.rgb = navy

box2 = slide.shapes.add_shape(1, Inches(4.8), Inches(1.5), Inches(3.7), Inches(4.5))
box2.fill.solid(); box2.fill.fore_color.rgb = white
box2.line.color.rgb = blue
box2.line.width = Pt(2)
box2_tf = box2.text_frame
box2_tf.word_wrap = True
add_paragraphs(
    box2_tf,
    [
        'Forecast & Risk Model',
        'A prediction model estimated next-day demand and classified SKUs as low, medium, or high risk.',
        'Risk scoring used cover ratio, projected demand, and stockout exposure.',
    ],
    font_size=17,
    color=text,
    space_after=8,
)
box2_tf.paragraphs[0].font.bold = True
box2_tf.paragraphs[0].font.size = Pt(22)
box2_tf.paragraphs[0].font.color.rgb = navy

box3 = slide.shapes.add_shape(1, Inches(9.0), Inches(1.5), Inches(3.5), Inches(4.5))
box3.fill.solid(); box3.fill.fore_color.rgb = navy
box3.line.color.rgb = navy
box3_tf = box3.text_frame
box3_tf.word_wrap = True
add_paragraphs(
    box3_tf,
    [
        'Decision Support',
        'A dashboard and AI assistant translated the output into understandable actions for managers.',
    ],
    font_size=17,
    color=white,
    space_after=8,
)
box3_tf.paragraphs[0].font.bold = True
box3_tf.paragraphs[0].font.size = Pt(22)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = white
add_title(slide, '3. What the Solution Delivers')

features = [
    ('SKU Risk Dashboard', 'View all SKUs and their current risk state in one place.'),
    ('Explainable Insights', 'See the actual stock, lead time, and demand drivers for each SKU.'),
    ('Natural Language Q&A', 'Ask business questions in plain English and get grounded answers.'),
    ('Operational Prioritization', 'Focus replenishment decisions on the SKUs with the highest risk first.'),
]
positions = [(0.8, 1.5, 2.8, 2.7), (3.9, 1.5, 2.8, 2.7), (7.0, 1.5, 2.8, 2.7), (10.1, 1.5, 2.8, 2.7)]
for (title, desc), (x, y, w, h) in zip(features, positions):
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = light
    box.line.color.rgb = blue
    box.line.width = Pt(2)
    tf = box.text_frame
    tf.word_wrap = True
    add_paragraphs(tf, [title, desc], font_size=13, color=text, space_after=8)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.color.rgb = navy

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = navy
add_title(slide, '4. Business Outcome')

box = slide.shapes.add_shape(1, Inches(1.0), Inches(1.8), Inches(11.2), Inches(3.8))
box.fill.solid(); box.fill.fore_color.rgb = white
box.line.color.rgb = white
box.line.width = Pt(2)
box_tf = box.text_frame
box_tf.word_wrap = True
add_paragraphs(
    box_tf,
    [
        'This solution helps supply chain teams:',
        '• Reduce stockout and overstock risk through faster, data-driven decisions',
        '• Prioritize inventory action based on actual SKU risk instead of intuition',
        '• Improve operational confidence with explainable model outputs and AI guidance',
    ],
    font_size=20,
    color=text,
    space_after=10,
)
box_tf.paragraphs[0].font.bold = True
box_tf.paragraphs[0].font.size = Pt(22)
box_tf.paragraphs[0].font.color.rgb = navy

final_line = slide.shapes.add_textbox(Inches(1.2), Inches(6.1), Inches(10.8), Inches(0.6))
ft = final_line.text_frame
p = ft.paragraphs[0]
p.text = 'Outcome: Better inventory visibility, smarter replenishment decisions, and reduced operational risk.'
p.font.size = Pt(20); p.font.bold = True; p.font.color.rgb = white

pptx_path = 'FMN_Supply_Chain_Solution_Presentation.pptx'
prs.save(pptx_path)
print(f'Created presentation: {pptx_path}')
